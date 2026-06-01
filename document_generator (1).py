import io
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from jinja2 import Template
from weasyprint import HTML
from llm_processor import OptimizedCV

class DocumentGenerator:
    TEMPLATES = {
        "Clásico": {"font": "Times New Roman", "body_size": 11, "name_size": 18, "color": "#000000", "layout": "single"},
        "Moderno": {"font": "Arial", "body_size": 10, "name_size": 20, "color": "#1f77b4", "layout": "single"},
        "Ivy League": {"font": "Georgia", "body_size": 10.5, "name_size": 22, "color": "#1a1a1a", "layout": "single"},
        "Elegante": {"font": "Garamond", "body_size": 11, "name_size": 24, "color": "#2c3e50", "layout": "single"},
        "Pulido": {"font": "Trebuchet MS", "body_size": 10, "name_size": 19, "color": "#34495e", "layout": "single"},
        "Contemporáneo": {"font": "Verdana", "body_size": 9.5, "name_size": 20, "color": "#16a085", "layout": "single"},
        "Creativo": {"font": "Tahoma", "body_size": 10, "name_size": 18, "color": "#e67e22", "layout": "single"},
        "Línea de tiempo": {"font": "Arial Narrow", "body_size": 10, "name_size": 18, "color": "#2980b9", "layout": "single"},
        "Estilizado": {"font": "Century Gothic", "body_size": 10, "name_size": 21, "color": "#8e44ad", "layout": "single"},
        "Compacto": {"font": "Calibri", "body_size": 9, "name_size": 14, "color": "#2c3e50", "layout": "single"},
        "Alto rendimiento": {"font": "Impact", "body_size": 11, "name_size": 26, "color": "#c0392b", "layout": "single"},
        "Minimalista": {"font": "Helvetica", "body_size": 11, "name_size": 16, "color": "#333333", "layout": "single"}
    }

    @staticmethod
    def generate_word(cv: OptimizedCV, template_name: str = "Clásico") -> bytes:
        config = DocumentGenerator.TEMPLATES.get(template_name, DocumentGenerator.TEMPLATES["Clásico"])
        doc = Document()
        margin = 0.5 if template_name == "Compacto" else 0.75
        for section in doc.sections:
            section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(margin)
        
        # Header
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(cv.full_name.upper())
        run.font.size = Pt(config["name_size"])
        run.font.bold = True
        run.font.name = config["font"]
        if template_name != "Clásico":
            run.font.color.rgb = RGBColor.from_string(config["color"].replace("#", ""))

        contact = " | ".join(filter(None, [cv.email, cv.phone, cv.location]))
        doc.add_paragraph(contact).alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Secciones
        def add_sec(title, content):
            if not content: return
            h = doc.add_paragraph(title.upper())
            h.runs[0].font.bold = True
            h.runs[0].font.size = Pt(config["body_size"] + 1)
            if template_name != "Clásico":
                h.runs[0].font.color.rgb = RGBColor.from_string(config["color"].replace("#", ""))
            doc.add_paragraph(content)

        add_sec("Resumen", cv.professional_summary)
        add_sec("Habilidades", ", ".join(cv.key_skills))

        if cv.experience:
            h = doc.add_paragraph("EXPERIENCIA")
            h.runs[0].font.bold = True
            for exp in cv.experience:
                p = doc.add_paragraph()
                p.add_run(f"{exp.job_title} | {exp.company}").font.bold = True
                doc.add_paragraph(exp.duration).runs[0].font.italic = True
                for ach in exp.achievements:
                    doc.add_paragraph(ach, style='List Bullet')

        if cv.education:
            h = doc.add_paragraph("EDUCACIÓN")
            h.runs[0].font.bold = True
            for edu in cv.education:
                doc.add_paragraph(f"{edu.degree} | {edu.institution} ({edu.year})")

        buffer = io.BytesIO()
        doc.save(buffer)
        return buffer.getvalue()

    @staticmethod
    def generate_pdf(cv: OptimizedCV, template_name: str = "Clásico") -> bytes:
        config = DocumentGenerator.TEMPLATES.get(template_name, DocumentGenerator.TEMPLATES["Clásico"])
        
        html_template = """
        <html>
        <head>
            <style>
                @page { margin: 0.6in; size: letter; }
                body { font-family: '{{font}}', sans-serif; font-size: {{body_size}}pt; line-height: 1.4; color: #333; }
                .header { text-align: center; border-bottom: 2px solid {{color}}; padding-bottom: 10px; margin-bottom: 20px; }
                h1 { font-size: {{name_size}}pt; color: {{color}}; margin: 0; text-transform: uppercase; }
                .contact { font-size: 9pt; margin-top: 5px; }
                h2 { font-size: {{body_size + 2}}pt; color: {{color}}; border-bottom: 1px solid #ddd; margin-top: 20px; text-transform: uppercase; }
                .job-item { margin-bottom: 15px; }
                .job-header { font-weight: bold; font-size: {{body_size + 1}}pt; display: flex; justify-content: space-between; }
                .company { color: {{color}}; }
                .duration { font-style: italic; color: #666; font-size: 9pt; }
                ul { margin-top: 5px; padding-left: 20px; }
                li { margin-bottom: 3px; }
                .skills-box { background: #f9f9f9; padding: 10px; border-radius: 5px; border-left: 4px solid {{color}}; }
            </style>
        </head>
        <body>
            <div class="header">
                <h1>{{cv.full_name}}</h1>
                <div class="contact">{{cv.email}} | {{cv.phone}} | {{cv.location}}</div>
            </div>
            
            <h2>Perfil Profesional</h2>
            <p>{{cv.professional_summary}}</p>
            
            <h2>Habilidades Técnicas</h2>
            <div class="skills-box">{{cv.key_skills|join(' • ')}}</div>
            
            <h2>Experiencia Profesional</h2>
            {% for exp in cv.experience %}
                <div class="job-item">
                    <div class="job-header">
                        <span>{{exp.job_title}} <span class="company">@ {{exp.company}}</span></span>
                        <span class="duration">{{exp.duration}}</span>
                    </div>
                    <ul>{% for ach in exp.achievements %}<li>{{ach}}</li>{% endfor %}</ul>
                </div>
            {% endfor %}
            
            <h2>Formación Académica</h2>
            {% for edu in cv.education %}
                <div class="job-item">
                    <div class="job-header">
                        <span>{{edu.degree}}</span>
                        <span class="duration">{{edu.year}}</span>
                    </div>
                    <div>{{edu.institution}}</div>
                </div>
            {% endfor %}
        </body>
        </html>
        """
        html = Template(html_template).render(cv=cv, **config)
        buffer = io.BytesIO()
        HTML(string=html).write_pdf(buffer)
        return buffer.getvalue()

    @staticmethod
    def generate_cl_word(content: str, name: str, template_name: str = "Clásico") -> bytes:
        return DocumentGenerator.generate_word(OptimizedCV(full_name=name, email="", professional_summary=content, key_skills=[], experience=[], education=[]), template_name)

    @staticmethod
    def generate_cl_pdf(content: str, name: str, template_name: str = "Clásico") -> bytes:
        config = DocumentGenerator.TEMPLATES.get(template_name, DocumentGenerator.TEMPLATES["Clásico"])
        html = f"""<html><body style="font-family: {config['font']}; padding: 40px; line-height: 1.6;">
            <h1 style="color: {config['color']}; border-bottom: 2px solid {config['color']}; padding-bottom: 10px;">{name}</h1>
            <div style="white-space: pre-wrap; margin-top: 30px; font-size: {config['body_size']}pt;">{content}</div>
        </body></html>"""
        buffer = io.BytesIO()
        HTML(string=html).write_pdf(buffer)
        return buffer.getvalue()
