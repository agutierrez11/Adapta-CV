import io
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from jinja2 import Template
from weasyprint import HTML
from llm_processor import OptimizedCV

class DocumentGenerator:
    TEMPLATES = {
        "Clásico": {"font": "Times New Roman", "body_size": 11, "name_size": 18, "color": "#000000"},
        "Moderno": {"font": "Arial", "body_size": 10, "name_size": 20, "color": "#1f77b4"},
        "Minimalista": {"font": "Calibri", "body_size": 11, "name_size": 16, "color": "#333333"}
    }

    @staticmethod
    def generate_word(cv: OptimizedCV, template_name: str = "Clásico") -> bytes:
        config = DocumentGenerator.TEMPLATES.get(template_name, DocumentGenerator.TEMPLATES["Clásico"])
        doc = Document()
        for section in doc.sections:
            section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(0.75)
        
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(cv.full_name.upper())
        run.font.size = Pt(config["name_size"])
        run.font.bold = True
        run.font.name = config["font"]

        contact = " | ".join(filter(None, [cv.email, cv.phone, cv.location]))
        doc.add_paragraph(contact).alignment = WD_ALIGN_PARAGRAPH.CENTER

        doc.add_paragraph("RESUMEN PROFESIONAL").runs[0].font.bold = True
        doc.add_paragraph(cv.professional_summary)

        doc.add_paragraph("HABILIDADES").runs[0].font.bold = True
        doc.add_paragraph(", ".join(cv.key_skills))

        doc.add_paragraph("EXPERIENCIA").runs[0].font.bold = True
        for exp in cv.experience:
            p = doc.add_paragraph()
            p.add_run(f"{exp.job_title} | {exp.company}").font.bold = True
            doc.add_paragraph(exp.duration).runs[0].font.italic = True
            for ach in exp.achievements:
                doc.add_paragraph(ach, style='List Bullet')

        doc.add_paragraph("EDUCACIÓN").runs[0].font.bold = True
        for edu in cv.education:
            doc.add_paragraph(f"{edu.degree} | {edu.institution} ({edu.year})")

        buffer = io.BytesIO()
        doc.save(buffer)
        return buffer.getvalue()

    @staticmethod
    def generate_pdf(cv: OptimizedCV, template_name: str = "Clásico") -> bytes:
        config = DocumentGenerator.TEMPLATES.get(template_name, DocumentGenerator.TEMPLATES["Clásico"])
        html_template = """
        <html><body style="font-family: {{font}}; font-size: {{body_size}}pt;">
            <h1 style="text-align: center; color: {{color}};">{{cv.full_name}}</h1>
            <p style="text-align: center;">{{cv.email}} | {{cv.phone}} | {{cv.location}}</p>
            <h2 style="color: {{color}};">RESUMEN</h2><p>{{cv.professional_summary}}</p>
            <h2 style="color: {{color}};">HABILIDADES</h2><p>{{cv.key_skills|join(', ')}}</p>
            <h2 style="color: {{color}};">EXPERIENCIA</h2>
            {% for exp in cv.experience %}
                <b>{{exp.job_title}} | {{exp.company}}</b><br><i>{{exp.duration}}</i>
                <ul>{% for ach in exp.achievements %}<li>{{ach}}</li>{% endfor %}</ul>
            {% endfor %}
            <h2 style="color: {{color}};">EDUCACIÓN</h2>
            {% for edu in cv.education %}<div><b>{{edu.degree}}</b> | {{edu.institution}} ({{edu.year}})</div>{% endfor %}
        </body></html>
        """
        html = Template(html_template).render(cv=cv, **config)
        buffer = io.BytesIO()
        HTML(string=html).write_pdf(buffer)
        return buffer.getvalue()

    @staticmethod
    def generate_cl_word(content: str, name: str, template_name: str = "Clásico") -> bytes:
        config = DocumentGenerator.TEMPLATES.get(template_name, DocumentGenerator.TEMPLATES["Clásico"])
        doc = Document()
        p = doc.add_paragraph(name.upper())
        p.runs[0].font.size = Pt(config["name_size"])
        p.runs[0].font.bold = True
        doc.add_paragraph(content)
        buffer = io.BytesIO()
        doc.save(buffer)
        return buffer.getvalue()

    @staticmethod
    def generate_cl_pdf(content: str, name: str, template_name: str = "Clásico") -> bytes:
        config = DocumentGenerator.TEMPLATES.get(template_name, DocumentGenerator.TEMPLATES["Clásico"])
        html = f"""<html><body style="font-family: {config['font']};">
            <h1 style="color: {config['color']};">{name}</h1>
            <div style="white-space: pre-wrap;">{content}</div>
        </body></html>"""
        buffer = io.BytesIO()
        HTML(string=html).write_pdf(buffer)
        return buffer.getvalue()
